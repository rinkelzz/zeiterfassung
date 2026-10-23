"""Monats- und Jahresberichte als PDF und Excel (Download und Mailversand)."""
import io
import sqlite3
from datetime import date, timedelta
from typing import Dict, List, Optional, Tuple

from . import csvsafe
from .pdf import PDF, fit
from .reports import (
    ABSENCE_KINDS,
    BREAK_MODES,
    break_rule,
    compute_report,
    schedule_on,
    tracking_start,
    vacation_account,
)
from .timeutil import fmt_duration, now_utc, to_local, today_local

MONTHS_DE = ["Januar", "Februar", "März", "April", "Mai", "Juni", "Juli", "August",
             "September", "Oktober", "November", "Dezember"]
WD_DE = ["Mo", "Di", "Mi", "Do", "Fr", "Sa", "So"]


# --- Perioden ----------------------------------------------------------------

def parse_period(period: str) -> Tuple[str, date, date]:
    """'2026-09' → Monat, '2026' → Jahr. Liefert (art, von, bis)."""
    p = (period or "").strip()
    try:
        if len(p) == 4:
            y = int(p)
            if 1990 <= y <= 2100:
                return "year", date(y, 1, 1), date(y, 12, 31)
        elif len(p) == 7 and p[4] == "-":
            y, m = int(p[:4]), int(p[5:])
            if 1990 <= y <= 2100 and 1 <= m <= 12:
                last = (date(y + (m == 12), m % 12 + 1, 1) - timedelta(days=1))
                return "month", date(y, m, 1), last
    except ValueError:
        pass
    raise ValueError("Zeitraum als JJJJ-MM (Monat) oder JJJJ (Jahr) angeben")


def period_title(period: str) -> str:
    kind, d_from, _ = parse_period(period)
    return f"{MONTHS_DE[d_from.month - 1]} {d_from.year}" if kind == "month" else f"Jahr {d_from.year}"


def file_name(user: sqlite3.Row, period: str, ext: str) -> str:
    kind, _, _ = parse_period(period)
    prefix = "Arbeitszeitnachweis" if kind == "month" else "Jahresuebersicht"
    return f"{prefix}_{user['username']}_{period}.{ext}"


# --- Daten ---------------------------------------------------------------------

def balance_until(conn: sqlite3.Connection, user: sqlite3.Row, d: date) -> Dict[str, int]:
    """Kontostände (inkl. Übertrag) am Ende des Tages d – ab dem letzten Monatsabschluss gerechnet."""
    from .reports import balance_at
    return balance_at(conn, user, d)


def _absence_counts(days: List[Dict]) -> Dict[str, float]:
    out = {}
    for k in ABSENCE_KINDS:
        n = sum((0.5 if x["half_day"] else 1) for x in days
                if x["absence"] == k and x["workday"] and not x["before_start"]
                and (k == "feiertag" or not x["holiday"]))
        if n:
            out[k] = n
    return out


def _schedule_text(conn, user, d: date) -> str:
    s = schedule_on(conn, user["id"], d)
    hours = s["weekly_hours"] if s else user["weekly_hours"]
    days = [int(x) for x in ((s["workdays"] if s else user["workdays"]) or "").split(",") if x]
    if days == [1, 2, 3, 4, 5]:
        dl = "Mo–Fr"
    else:
        dl = ", ".join(WD_DE[x - 1] for x in sorted(days))
    return f"{_num(hours)} Std./Woche ({dl})"


def build_data(conn: sqlite3.Connection, user: sqlite3.Row, period: str) -> Dict:
    kind, d_from, d_to = parse_period(period)
    rep = compute_report(conn, user, d_from, d_to)
    start_bal = balance_until(conn, user, d_from - timedelta(days=1))
    end_bal = balance_until(conn, user, d_to)
    data = {
        "kind": kind,
        "period": period,
        "title": period_title(period),
        "user": user,
        "from": d_from,
        "to": d_to,
        "report": rep,
        "start_balance": start_bal,
        "end_balance": end_bal,
        "end_is_today": d_to >= today_local(),
        "vacation": vacation_account(conn, user, d_to.year),
        "absences": _absence_counts(rep["days"]),
        "schedule": _schedule_text(conn, user, min(d_to, today_local()) if d_from <= today_local() else d_from),
        "break_rule": break_rule(conn),
        "created": to_local(now_utc()),
        "closed_at": None,
    }
    if kind == "month":
        row = conn.execute("SELECT closed_at FROM month_closings WHERE user_id = ? AND month = ?",
                           (user["id"], period)).fetchone()
        if row:
            from .timeutil import parse_utc
            data["closed_at"] = to_local(parse_utc(row["closed_at"]))
    if kind == "year":
        months = []
        for m in range(1, 13):
            days = [x for x in rep["days"] if int(x["date"][5:7]) == m]
            months.append({
                "month": m,
                "name": MONTHS_DE[m - 1],
                "target_s": sum(x["target_s"] for x in days),
                "worked_s": sum(x["worked_s"] for x in days),
                "diff_s": sum(x["diff_s"] for x in days),
                "gleitzeit_s": sum(x["diff_gleitzeit_s"] for x in days),
                "ueberstunden_s": sum(x["diff_ueberstunden_s"] for x in days),
                "travel_min": sum(x["travel_min"] for x in days),
                "absences": _absence_counts(days),
            })
        data["months"] = months
    return data


# --- Formatierung ----------------------------------------------------------------

def _num(v: float) -> str:
    v = round(float(v or 0), 2)
    return (f"{v:g}").replace(".", ",")


def _dur(s: float, signed: bool = False) -> str:
    txt = fmt_duration(s)
    return ("+" + txt) if signed and s > 0 else txt


def _day_note(d: Dict) -> str:
    parts = []
    if d["absence_label"]:
        parts.append(d["absence_label"])
    elif d["holiday"]:
        parts.append(d["holiday"])
    if d["overtime_day"]:
        parts.append("Ü-Tag")
    if d["travel_min"]:
        parts.append(f"Wegezeit {d['travel_min']} min")
    if d["segments"] > 1:
        parts.append(f"{d['segments']} Buchungen")
    if d["running"]:
        parts.append("läuft")
    if d.get("needs_check"):
        parts.insert(0, "autom. beendet – prüfen")
    for w in reversed(d.get("warnings") or []):
        parts.insert(0, "ArbZG: " + w["text"])
    if d["before_start"]:
        parts.append("vor Erfassungsbeginn")
    return ", ".join(parts)


def _counts(d: Dict) -> bool:
    return not d["before_start"] and (d["target_s"] or d["worked_s"] or d["diff_s"])


# --- PDF -------------------------------------------------------------------------

GREY = (0.45, 0.45, 0.45)
LIGHT = (0.94, 0.95, 0.97)
HEAD = (0.86, 0.89, 0.94)
NEG = (0.75, 0.1, 0.1)
POS = (0.05, 0.45, 0.2)
L, R = 40, 555  # linker/rechter Rand (A4 = 595 pt breit)


def _pdf_header(pdf: PDF, data: Dict, heading: str) -> float:
    u = data["user"]
    pdf.text(L, 58, heading, 17, True)
    pdf.text(R, 58, data["title"], 13, True, align="right")
    pdf.text(L, 78, f"{u['display_name']}  ({u['username']})", 10)
    pdf.text(R, 78, f"{data['from']:%d.%m.%Y} – {data['to']:%d.%m.%Y}", 10, align="right")
    rule = data["break_rule"]
    brk = BREAK_MODES.get(rule["mode"], "")
    if rule["mode"] == "fixed":
        brk = f"fest {rule['minutes']} min pro Arbeitstag"
    pdf.text(L, 92, f"Sollarbeitszeit: {data['schedule']}   ·   Pausenabzug: {brk}", 8, color=GREY)
    if data.get("closed_at"):
        pdf.text(R, 92, f"Monat abgeschlossen am {data['closed_at']:%d.%m.%Y}", 8, True, "right", (0.14, 0.25, 0.48))
    pdf.line(L, 100, R, 100, 1, (0.3, 0.3, 0.3))
    return 116


def _pdf_footer(pdf: PDF, data: Dict, signatures: bool = True) -> None:
    if signatures:
        y = 770
        pdf.line(L, y, L + 220, y, 0.6, (0.2, 0.2, 0.2))
        pdf.line(R - 220, y, R, y, 0.6, (0.2, 0.2, 0.2))
        pdf.text(L, y + 11, "Datum, Unterschrift Mitarbeiter/in", 7.5, color=GREY)
        pdf.text(R - 220, y + 11, "Datum, Unterschrift Vorgesetzte/r", 7.5, color=GREY)
    pdf.text(L, 815, f"Erstellt am {data['created']:%d.%m.%Y um %H:%M} Uhr · Zeiterfassung", 7, color=GREY)


def _signed_color(s: float):
    return NEG if s < 0 else POS if s > 0 else (0, 0, 0)


def _accounts_block(pdf: PDF, data: Dict, y: float, label_start: str, label_end: str) -> float:
    sb, eb, t = data["start_balance"], data["end_balance"], data["report"]["totals"]
    pdf.text(L, y, "Konten", 10.5, True)
    y += 6
    cols = [(L, "left", ""), (L + 270, "right", "Gleitzeit"), (L + 370, "right", "Überstunden")]
    pdf.rect(L, y, 380, 14, HEAD)
    for x, al, title in cols:
        pdf.text(x + (4 if al == "left" else 0), y + 10, title, 8.5, True, align=al)
    y += 14
    rows = [
        (label_start, sb["gleitzeit_s"], sb["ueberstunden_s"], False),
        ("Veränderung im Zeitraum", t["gleitzeit_s"], t["ueberstunden_s"], True),
        (label_end, eb["gleitzeit_s"], eb["ueberstunden_s"], False),
    ]
    for i, (label, g, u, signed) in enumerate(rows):
        bold = i == 2
        pdf.text(L + 4, y + 10, label, 8.5, bold)
        pdf.text(L + 270, y + 10, _dur(g, True), 8.5, bold, "right", _signed_color(g))
        pdf.text(L + 370, y + 10, _dur(u, True), 8.5, bold, "right", _signed_color(u))
        y += 14
        pdf.line(L, y, L + 380, y, 0.3, (0.8, 0.8, 0.8))
    return y + 18


def _vacation_block(pdf: PDF, data: Dict, y: float) -> float:
    vac = data["vacation"]
    absn = data["absences"]
    if vac:
        pdf.text(L, y, f"Urlaub {vac['year']}", 10.5, True)
        y += 14
        pdf.text(L, y, f"Anspruch {_num(vac['entitlement'])}  ·  Übertrag {_num(vac['carryover'])}  ·  "
                       f"genommen {_num(vac['taken'])}  ·  geplant {_num(vac['planned'])}  ·  "
                       f"Rest {_num(vac['remaining'])} Tage", 8.5)
        y += 18
    if absn:
        pdf.text(L, y, "Abwesenheiten im Zeitraum", 10.5, True)
        y += 14
        pdf.text(L, y, fit("  ·  ".join(f"{ABSENCE_KINDS[k]} {_num(v)} Tag{'e' if v != 1 else ''}" for k, v in absn.items()), R - L, 8.5), 8.5)
        y += 18
    return y


def month_pdf(data: Dict) -> bytes:
    pdf = PDF(f"Arbeitszeitnachweis {data['title']} – {data['user']['display_name']}")
    y = _pdf_header(pdf, data, "Arbeitszeitnachweis")
    # Spalten: (Titel, Breite, Ausrichtung)
    cols = [("Datum", 50, "left"), ("Tag", 22, "left"), ("Kommen", 40, "left"), ("Gehen", 40, "left"),
            ("Pause", 34, "right"), ("Ist", 40, "right"), ("Soll", 40, "right"), ("Differenz", 48, "right"),
            ("Konto", 34, "left"), ("Bemerkung", 167, "left")]
    xs, x = [], L
    for _, w, _a in cols:
        xs.append(x)
        x += w

    def cell_x(i):
        return xs[i] + cols[i][1] - 4 if cols[i][2] == "right" else xs[i] + 3

    rh = 13.2
    pdf.rect(L, y, R - L, rh + 1, HEAD)
    for i, (title, _, al) in enumerate(cols):
        pdf.text(cell_x(i), y + 9.5, title, 7.8, True, align=al)
    y += rh + 1
    for d in data["report"]["days"]:
        dd = date.fromisoformat(d["date"])
        off = not d["workday"] or d["holiday"]
        if off:
            pdf.rect(L, y, R - L, rh, LIGHT)
        counts = _counts(d)
        konto = ""
        if counts and d["diff_s"]:
            konto = "Ü" if d["diff_ueberstunden_s"] else "GZ"
        vals = [
            f"{dd:%d.%m.}", d["weekday"], d["first_in"] or "",
            ("läuft" if d["running"] else d["last_out"]) or "",
            f"{d['break_deducted_s'] // 60} min" if d["break_deducted_s"] else "",
            d["worked"] if d["worked_s"] else "",
            d["target"] if d["target_s"] else "",
            _dur(d["diff_s"], True) if counts and d["diff_s"] else "",
            konto,
            fit(_day_note(d), cols[9][1] - 6, 7.5),
        ]
        for i, v in enumerate(vals):
            color = _signed_color(d["diff_s"]) if i == 7 else (GREY if i == 9 else (0, 0, 0))
            pdf.text(cell_x(i), y + 9.4, v, 7.5 if i == 9 else 8, align=cols[i][2], color=color)
        y += rh
        pdf.line(L, y, R, y, 0.25, (0.85, 0.85, 0.85))
    t = data["report"]["totals"]
    pdf.line(L, y, R, y, 0.8, (0.3, 0.3, 0.3))
    y += 1
    pdf.text(cell_x(0), y + 10, "Summe", 8.5, True)
    pdf.text(cell_x(4), y + 10, f"{t['break_deducted_s'] // 60} min" if t["break_deducted_s"] else "", 8, True, "right")
    pdf.text(cell_x(5), y + 10, t["worked"], 8.5, True, "right")
    pdf.text(cell_x(6), y + 10, t["target"], 8.5, True, "right")
    pdf.text(cell_x(7), y + 10, _dur(t["diff_s"], True), 8.5, True, "right", _signed_color(t["diff_s"]))
    travel = sum(d["travel_min"] for d in data["report"]["days"])
    if travel:
        pdf.text(cell_x(9), y + 10, f"davon Wegezeit {fmt_duration(travel * 60)} h", 7.5, color=GREY)
    if t.get("warnings"):
        n = t["warnings"]
        pdf.text(L, y + 26, f"{n} Hinweis{'e' if n != 1 else ''} nach Arbeitszeitgesetz (über 10 h am Tag oder unter "
                            "11 h Ruhezeit) – siehe Bemerkung.", 8, color=NEG)
        y += 14
    y += 30

    end_label = "Stand heute" if data["end_is_today"] else f"Stand {data['to']:%d.%m.%Y}"
    y = _accounts_block(pdf, data, y, f"Stand {data['from'] - timedelta(days=1):%d.%m.%Y} (inkl. Übertrag)", end_label)
    _vacation_block(pdf, data, y)
    _pdf_footer(pdf, data)
    return pdf.output()


def year_pdf(data: Dict) -> bytes:
    pdf = PDF(f"Jahresübersicht {data['from'].year} – {data['user']['display_name']}")
    y = _pdf_header(pdf, data, "Jahresübersicht")
    cols = [("Monat", 78, "left"), ("Soll", 52, "right"), ("Ist", 52, "right"), ("Differenz", 56, "right"),
            ("Gleitzeit", 56, "right"), ("Überstunden", 62, "right"), ("Urlaub", 40, "right"),
            ("Krank", 40, "right"), ("Sonstige", 79, "right")]
    xs, x = [], L
    for _, w, _a in cols:
        xs.append(x)
        x += w

    def cell_x(i):
        return xs[i] + cols[i][1] - 4 if cols[i][2] == "right" else xs[i] + 4

    rh = 18
    pdf.rect(L, y, R - L, rh, HEAD)
    for i, (title, _, al) in enumerate(cols):
        pdf.text(cell_x(i), y + 12, title, 8.5, True, align=al)
    y += rh
    tot_abs: Dict[str, float] = {}
    for i, m in enumerate(data["months"]):
        if i % 2:
            pdf.rect(L, y, R - L, rh, LIGHT)
        a = m["absences"]
        for k, v in a.items():
            tot_abs[k] = tot_abs.get(k, 0) + v
        other = sum(v for k, v in a.items() if k not in ("urlaub", "krank", "feiertag"))
        vals = [m["name"], fmt_duration(m["target_s"]) if m["target_s"] else "",
                fmt_duration(m["worked_s"]) if m["worked_s"] else "",
                _dur(m["diff_s"], True) if m["diff_s"] else "",
                _dur(m["gleitzeit_s"], True) if m["gleitzeit_s"] else "",
                _dur(m["ueberstunden_s"], True) if m["ueberstunden_s"] else "",
                _num(a.get("urlaub", 0)) if a.get("urlaub") else "",
                _num(a.get("krank", 0)) if a.get("krank") else "",
                _num(other) if other else ""]
        for j, v in enumerate(vals):
            color = _signed_color(m[["diff_s", "gleitzeit_s", "ueberstunden_s"][j - 3]]) if j in (3, 4, 5) else (0, 0, 0)
            pdf.text(cell_x(j), y + 12, v, 9, align=cols[j][2], color=color)
        y += rh
    t = data["report"]["totals"]
    pdf.line(L, y, R, y, 0.8, (0.3, 0.3, 0.3))
    other = sum(v for k, v in tot_abs.items() if k not in ("urlaub", "krank", "feiertag"))
    vals = ["Summe", t["target"], t["worked"], _dur(t["diff_s"], True), _dur(t["gleitzeit_s"], True),
            _dur(t["ueberstunden_s"], True)] + [_num(v) if v else "" for v in (tot_abs.get("urlaub", 0), tot_abs.get("krank", 0), other)]
    for j, v in enumerate(vals):
        color = _signed_color([t["diff_s"], t["gleitzeit_s"], t["ueberstunden_s"]][j - 3]) if j in (3, 4, 5) else (0, 0, 0)
        pdf.text(cell_x(j), y + 13, v, 9, True, cols[j][2], color)
    y += 34
    pdf.text(L, y, "Urlaub, Krank und Sonstige in Arbeitstagen. Sonstige: Bildungsurlaub, Gleittag, "
                   "Mehrarbeit-Entnahme, Sonderurlaub.", 7.5, color=GREY)
    y += 26
    end_label = "Stand heute" if data["end_is_today"] else f"Stand {data['to']:%d.%m.%Y}"
    y = _accounts_block(pdf, data, y, f"Stand {data['from'] - timedelta(days=1):%d.%m.%Y} (inkl. Übertrag)", end_label)
    _vacation_block(pdf, data, y)
    _pdf_footer(pdf, data)
    return pdf.output()


# --- Excel -------------------------------------------------------------------------

def _hours(s: float) -> float:
    return round((s or 0) / 3600, 6)  # exakt speichern, angezeigt wird mit 2 Nachkommastellen


def _xlsx_month_sheet(ws, data: Dict, days: List[Dict], title_line: str) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DCE3EF")
    off_fill = PatternFill("solid", fgColor="F0F2F7")
    hfmt = '+0.00;-0.00;0.00'
    u = data["user"]
    ws.append([title_line])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"{u['display_name']} ({u['username']})", "", "", "", "", "Sollarbeitszeit:", data["schedule"]])
    csvsafe.no_formulas(ws)
    ws.append([])
    header = ["Datum", "Tag", "Kommen", "Gehen", "Pause (min)", "Ist (h)", "Soll (h)", "Differenz (h)",
              "davon Gleitzeit (h)", "davon Überstunden (h)", "Wegezeit (min)", "Abwesenheit / Feiertag", "Bemerkung"]
    ws.append(header)
    hr = ws.max_row
    for c in ws[hr]:
        c.font = bold
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
    first = hr + 1
    for d in days:
        counts = _counts(d)
        note = []
        if d["overtime_day"]:
            note.append("Ü-Tag")
        if d["running"]:
            note.append("läuft")
        if d.get("needs_check"):
            note.append("automatisch beendet – bitte prüfen")
        note += ["ArbZG: " + w["text"] for w in d.get("warnings") or []]
        if d["before_start"]:
            note.append("vor Erfassungsbeginn")
        ws.append([
            date.fromisoformat(d["date"]), d["weekday"], d["first_in"] or "",
            ("läuft" if d["running"] else d["last_out"]) or "",
            d["break_deducted_s"] // 60 if d["break_deducted_s"] else None,
            _hours(d["worked_s"]) if d["worked_s"] else None,
            _hours(d["target_s"]) if d["target_s"] else None,
            _hours(d["diff_s"]) if counts else None,
            _hours(d["diff_gleitzeit_s"]) if counts and d["diff_gleitzeit_s"] else None,
            _hours(d["diff_ueberstunden_s"]) if counts and d["diff_ueberstunden_s"] else None,
            d["travel_min"] or None,
            d["absence_label"] or d["holiday"] or "",
            ", ".join(note),
        ])
        r = ws.max_row
        ws.cell(r, 1).number_format = "DD.MM.YYYY"
        for col in (6, 7):
            ws.cell(r, col).number_format = "0.00"
        for col in (8, 9, 10):
            ws.cell(r, col).number_format = hfmt
        if not d["workday"] or d["holiday"]:
            for c in ws[r]:
                c.fill = off_fill
    last = ws.max_row
    ws.append(["Summe", "", "", "",
               f"=SUM(E{first}:E{last})", f"=SUM(F{first}:F{last})", f"=SUM(G{first}:G{last})",
               f"=SUM(H{first}:H{last})", f"=SUM(I{first}:I{last})", f"=SUM(J{first}:J{last})",
               f"=SUM(K{first}:K{last})"])
    r = ws.max_row
    for c in ws[r]:
        c.font = bold
    for col in (6, 7):
        ws.cell(r, col).number_format = "0.00"
    for col in (8, 9, 10):
        ws.cell(r, col).number_format = hfmt
    widths = [11, 5, 9, 9, 8, 8, 8, 10, 11, 12, 9, 24, 22]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.row_dimensions[hr].height = 30
    ws.freeze_panes = ws.cell(hr + 1, 1)


def _xlsx_summary(ws, data: Dict, row_offset_blank: bool = True) -> None:
    from openpyxl.styles import Font

    bold = Font(bold=True)
    hfmt = '+0.00;-0.00;0.00'
    sb, eb, t = data["start_balance"], data["end_balance"], data["report"]["totals"]
    if row_offset_blank:
        ws.append([])
    ws.append(["Konten (Stunden)", "", "", "", "", "Gleitzeit", "Überstunden"])
    ws.cell(ws.max_row, 1).font = bold
    end_label = "Stand heute" if data["end_is_today"] else f"Stand {data['to']:%d.%m.%Y}"
    for label, g, u in ((f"Stand {data['from'] - timedelta(days=1):%d.%m.%Y} (inkl. Übertrag)", sb["gleitzeit_s"], sb["ueberstunden_s"]),
                        ("Veränderung im Zeitraum", t["gleitzeit_s"], t["ueberstunden_s"]),
                        (end_label, eb["gleitzeit_s"], eb["ueberstunden_s"])):
        ws.append([label, "", "", "", "", _hours(g), _hours(u)])
        ws.cell(ws.max_row, 6).number_format = hfmt
        ws.cell(ws.max_row, 7).number_format = hfmt
    vac = data["vacation"]
    if vac:
        ws.append([])
        ws.append([f"Urlaub {vac['year']} (Tage)", "", "", "", "", "Anspruch", "Übertrag", "genommen", "geplant", "Rest"])
        ws.cell(ws.max_row, 1).font = bold
        ws.append(["", "", "", "", "", vac["entitlement"], vac["carryover"], vac["taken"], vac["planned"], vac["remaining"]])
    if data["absences"]:
        ws.append([])
        ws.append(["Abwesenheiten im Zeitraum (Arbeitstage)"])
        ws.cell(ws.max_row, 1).font = bold
        for k, v in data["absences"].items():
            ws.append([ABSENCE_KINDS[k], "", "", "", "", v])


def month_xlsx(data: Dict) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = data["title"][:31]
    _xlsx_month_sheet(ws, data, data["report"]["days"], f"Arbeitszeitnachweis {data['title']}")
    _xlsx_summary(ws, data)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def year_xlsx(data: Dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    bold = Font(bold=True)
    hfmt = '+0.00;-0.00;0.00'
    wb = Workbook()
    ws = wb.active
    ws.title = "Jahresübersicht"
    u = data["user"]
    ws.append([f"Jahresübersicht {data['from'].year}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"{u['display_name']} ({u['username']})", "", "", "", "", "Sollarbeitszeit:", data["schedule"]])
    csvsafe.no_formulas(ws)
    ws.append([])
    kinds = list(ABSENCE_KINDS)
    ws.append(["Monat", "Soll (h)", "Ist (h)", "Differenz (h)", "Gleitzeit (h)", "Überstunden (h)", "Wegezeit (min)"]
              + [ABSENCE_KINDS[k] + " (Tage)" for k in kinds])
    hr = ws.max_row
    for c in ws[hr]:
        c.font = bold
        c.fill = PatternFill("solid", fgColor="DCE3EF")
    for m in data["months"]:
        ws.append([m["name"], _hours(m["target_s"]), _hours(m["worked_s"]), _hours(m["diff_s"]),
                   _hours(m["gleitzeit_s"]), _hours(m["ueberstunden_s"]), m["travel_min"]]
                  + [m["absences"].get(k, 0) for k in kinds])
        r = ws.max_row
        for col in (2, 3):
            ws.cell(r, col).number_format = "0.00"
        for col in (4, 5, 6):
            ws.cell(r, col).number_format = hfmt
    first, last = hr + 1, ws.max_row
    ws.append(["Summe"] + [f"=SUM({chr(65 + c)}{first}:{chr(65 + c)}{last})" for c in range(1, 7 + len(kinds))])
    r = ws.max_row
    for c in ws[r]:
        c.font = bold
    for col in (2, 3):
        ws.cell(r, col).number_format = "0.00"
    for col in (4, 5, 6):
        ws.cell(r, col).number_format = hfmt
    ws.column_dimensions["A"].width = 30
    for i in range(1, 7 + len(kinds)):
        ws.column_dimensions[chr(65 + i)].width = 13
    ws.row_dimensions[hr].height = 30
    _xlsx_summary(ws, data)

    for m in range(1, 13):
        days = [x for x in data["report"]["days"] if int(x["date"][5:7]) == m]
        sheet = wb.create_sheet(MONTHS_DE[m - 1])
        _xlsx_month_sheet(sheet, data, days, f"{MONTHS_DE[m - 1]} {data['from'].year}")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- Einstieg ------------------------------------------------------------------------

def render(conn: sqlite3.Connection, user: sqlite3.Row, period: str, fmt: str) -> Tuple[bytes, str, str]:
    """Liefert (Inhalt, Dateiname, MIME-Typ)."""
    return render_data(build_data(conn, user, period), fmt)


def render_data(data: Dict, fmt: str) -> Tuple[bytes, str, str]:
    user, period = data["user"], data["period"]
    if fmt == "pdf":
        body = month_pdf(data) if data["kind"] == "month" else year_pdf(data)
        return body, file_name(user, period, "pdf"), "application/pdf"
    if fmt == "xlsx":
        body = month_xlsx(data) if data["kind"] == "month" else year_xlsx(data)
        return body, file_name(user, period, "xlsx"), \
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    raise ValueError("Format muss pdf oder xlsx sein")


def previous_month(d: date) -> str:
    first = d.replace(day=1) - timedelta(days=1)
    return f"{first.year}-{first.month:02d}"


def default_period(kind: str, today: Optional[date] = None) -> str:
    today = today or today_local()
    return f"{today.year}-{today.month:02d}" if kind == "month" else str(today.year)

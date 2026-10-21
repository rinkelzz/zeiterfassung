"""Monatsabschluss: der Reihe nach, Sperren, festgeschriebene Kontostände, Wieder öffnen."""
import os

import pytest
from fastapi.testclient import TestClient

from app.main import app


@pytest.fixture(scope="module")
def admin():
    with TestClient(app) as c:
        assert c.post("/api/v1/auth/login", json={"username": "admin", "password": os.environ["ZE_ADMIN_PASSWORD"]}).status_code == 200
        c.patch("/api/v1/settings", json={"break_mode": "none", "bundesland": "NW", "company_free_days": ""})
        yield c


@pytest.fixture(scope="module")
def uid(admin):
    r = admin.post("/api/v1/users", json={"username": "abschluss", "display_name": "Abby", "password": "geheim123",
                                         "tracking_start": "2026-03-01", "weekly_hours": 40, "carryover_minutes": 60})
    uid = r.json()["id"]
    # März: Mo 02.03. 10 h → +2 h; April: Mi 01.04. 9 h → +1 h
    for s, e in (("2026-03-02T07:00", "2026-03-02T17:00"), ("2026-04-01T07:00", "2026-04-01T16:00")):
        assert admin.post("/api/v1/entries", json={"user_id": uid, "start": s, "end": e}).status_code == 201
    return uid


def st(admin, uid, month):
    return admin.get("/api/v1/closings", params={"user_id": uid, "month": month}).json()


def test_order_and_close(admin, uid):
    assert st(admin, uid, "2026-03")["can_close"] is True
    assert st(admin, uid, "2026-04")["can_close"] is False
    r = admin.post("/api/v1/closings", json={"user_id": uid, "month": "2026-04"})
    assert r.status_code == 409 and "März 2026" in r.json()["detail"]
    r = admin.post("/api/v1/closings", json={"user_id": uid, "month": "2026-03"})
    assert r.status_code == 200, r.text
    snap = r.json()
    # Übertrag 1 h + 2 h am 02.03. − Soll aller übrigen Arbeitstage im März
    rep = admin.get("/api/v1/reports/user", params={"user_id": uid, "from": "2026-03-01", "to": "2026-03-31"}).json()
    assert snap["gleitzeit_s"] == 3600 + rep["totals"]["gleitzeit_s"]
    s = st(admin, uid, "2026-03")
    assert s["closed"] and s["can_reopen"] and s["next_month"] == "2026-04"
    assert admin.post("/api/v1/closings", json={"user_id": uid, "month": "2026-03"}).status_code == 409
    # Zukunft/laufender Monat geht nicht
    r = admin.post("/api/v1/closings", json={"user_id": uid, "month": "2099-01"})
    assert r.status_code == 409


def test_locks(admin, uid):
    r = admin.post("/api/v1/entries", json={"user_id": uid, "start": "2026-03-03T07:00", "end": "2026-03-03T15:00"})
    assert r.status_code == 409 and "März 2026 ist abgeschlossen" in r.json()["detail"]
    eid = admin.get("/api/v1/entries", params={"user_id": uid, "from": "2026-03-02", "to": "2026-03-02"}).json()[0]["id"]
    assert admin.patch(f"/api/v1/entries/{eid}", json={"end": "2026-03-02T16:00"}).status_code == 409
    assert admin.delete(f"/api/v1/entries/{eid}").status_code == 409
    # Buchung aus April darf nicht in den März verschoben werden
    aid = admin.get("/api/v1/entries", params={"user_id": uid, "from": "2026-04-01", "to": "2026-04-01"}).json()[0]["id"]
    assert admin.patch(f"/api/v1/entries/{aid}", json={"start": "2026-03-31T07:00", "end": "2026-03-31T16:00"}).status_code == 409
    assert admin.post("/api/v1/absences", json={"user_id": uid, "from_day": "2026-03-30", "to_day": "2026-04-02",
                                               "kind": "urlaub"}).status_code == 409
    assert admin.put("/api/v1/dayflags", json={"user_id": uid, "day": "2026-03-02", "overtime": True}).status_code == 409
    assert admin.delete("/api/v1/absences", params={"user_id": uid, "from": "2026-03-01", "to": "2026-03-31"}).status_code == 409
    assert admin.post(f"/api/v1/users/{uid}/schedules", json={"valid_from": "2026-03-15", "weekly_hours": 30,
                                                            "workdays": [1, 2, 3, 4, 5]}).status_code == 409
    r = admin.patch(f"/api/v1/users/{uid}", json={"carryover_minutes": 600})
    assert r.status_code == 409 and "gesperrt" in r.json()["detail"]
    # unveränderte Werte mitschicken ist ok (Speichern-Knopf in der Mitarbeitertabelle)
    assert admin.patch(f"/api/v1/users/{uid}", json={"carryover_minutes": 60, "tracking_start": "2026-03-01",
                                                    "vacation_days_per_year": 28}).status_code == 200
    # Import in den März wird abgewiesen
    csv_text = "datum;checkin;checkout;dauer_h;notiz\n2026-03-04;08:00;12:00;4;x\n2026-04-06;08:00;12:00;4;y\n"
    r = admin.post("/api/v1/import", json={"user_id": uid, "kind": "entries", "csv_text": csv_text, "commit": False}).json()
    assert [x["status"] for x in r["rows"]] == ["fehler", "neu"] and "abgeschlossen" in r["rows"][0]["message"]
    # April ist offen
    assert admin.post("/api/v1/absences", json={"user_id": uid, "from_day": "2026-04-02", "kind": "urlaub"}).status_code == 201


def test_balance_frozen(admin, uid):
    before = admin.get("/api/v1/reports/user", params={"user_id": uid, "from": "2026-04-01", "to": "2026-04-30"}).json()["account"]
    # Wochenstunden ändern wirkt nur auf offene Monate: neuer Verlaufseintrag ab 01.04.
    assert admin.patch(f"/api/v1/users/{uid}", json={"weekly_hours": 20}).status_code == 200
    sched = admin.get(f"/api/v1/users/{uid}/schedules").json()
    froms = [s["valid_from"] for s in (sched["items"] if isinstance(sched, dict) else sched)]
    assert "2026-04-01" in froms
    march = admin.get("/api/v1/reports/user", params={"user_id": uid, "from": "2026-03-02", "to": "2026-03-02"}).json()
    assert march["days"][0]["target"] == "8:00"                      # März unverändert 40 h/5
    snap = st(admin, uid, "2026-03")["closed"]["gleitzeit_s"]
    after = admin.get("/api/v1/reports/user", params={"user_id": uid, "from": "2026-04-01", "to": "2026-04-30"}).json()["account"]
    assert after["gleitzeit_s"] != before["gleitzeit_s"]               # April rechnet mit 4 h/Tag
    # PDF vom März zeigt den Abschluss und den festgeschriebenen Stand
    pdf = admin.get("/api/v1/reports/download", params={"user_id": uid, "period": "2026-03", "format": "pdf"})
    assert pdf.status_code == 200
    from openpyxl import load_workbook
    import io
    x = admin.get("/api/v1/reports/download", params={"user_id": uid, "period": "2026-04", "format": "xlsx"})
    rows = list(load_workbook(io.BytesIO(x.content)).active.iter_rows(values_only=True))
    start_row = [r for r in rows if r[0] and str(r[0]).startswith("Stand 31.03.2026")][0]
    assert abs(start_row[5] - snap / 3600) < 0.01


def test_reopen_and_close_all(admin, uid):
    assert admin.post("/api/v1/closings/reopen", json={"user_id": uid, "month": "2026-03", "reason": ""}).status_code == 400
    r = admin.post("/api/v1/closings/reopen", json={"user_id": uid, "month": "2026-03", "reason": "Nachtrag Arzttermin"})
    assert r.status_code == 200
    log = admin.get("/api/v1/audit").json()
    assert any(a["action"] == "month_reopened" and "Arzttermin" in a["details"] for a in log)
    assert admin.post("/api/v1/entries", json={"user_id": uid, "start": "2026-03-03T07:00", "end": "2026-03-03T15:00"}).status_code == 201
    r = admin.post("/api/v1/closings/all", json={"until": "2026-04"}).json()
    mine = [x for x in r["results"] if x["user_id"] == uid][0]
    assert mine["closed"] == ["2026-03", "2026-04"] and mine["error"] is None
    # nur der letzte abgeschlossene Monat lässt sich öffnen
    r = admin.post("/api/v1/closings/reopen", json={"user_id": uid, "month": "2026-03", "reason": "Test"})
    assert r.status_code == 409 and "April 2026" in r.json()["detail"]
    ov = {x["user_id"]: x for x in admin.get("/api/v1/closings/overview").json()}
    assert ov[uid]["last_closed"]["month"] == "2026-04" and ov[uid]["next_month"] == "2026-05"


def test_rights(admin, uid):
    with TestClient(app) as c:
        admin.post("/api/v1/users", json={"username": "nichtadmin", "display_name": "N", "password": "geheim123"})
        c.post("/api/v1/auth/login", json={"username": "nichtadmin", "password": "geheim123"})
        assert c.post("/api/v1/closings", json={"user_id": uid, "month": "2026-05"}).status_code == 403
        assert c.get("/api/v1/closings", params={"month": "2026-03"}).status_code == 200   # eigener Status
        assert c.get("/api/v1/closings", params={"user_id": uid}).status_code == 403

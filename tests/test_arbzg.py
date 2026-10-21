"""Hinweise nach Arbeitszeitgesetz: über 10 h am Tag (§ 3), unter 11 h Ruhezeit (§ 5)."""
import io
import os

import pytest
from fastapi.testclient import TestClient

from app.main import app


@pytest.fixture(scope="module")
def admin():
    with TestClient(app) as c:
        assert c.post("/api/v1/auth/login", json={"username": "admin", "password": os.environ["ZE_ADMIN_PASSWORD"]}).status_code == 200
        c.patch("/api/v1/settings", json={"break_mode": "none", "arbzg_warnings": True, "bundesland": "NW"})
        yield c


@pytest.fixture(scope="module")
def uid(admin):
    r = admin.post("/api/v1/users", json={"username": "arbzg", "display_name": "Arno", "password": "geheim123",
                                         "tracking_start": "2026-01-01"})
    uid = r.json()["id"]
    for s, e in (
        ("2026-06-01T06:00", "2026-06-01T17:00"),   # Mo: 11 h → über 10 h
        ("2026-06-02T06:00", "2026-06-02T11:00"),   # Di: 5 h (Ruhezeit 13 h) …
        ("2026-06-02T17:00", "2026-06-02T21:00"),   # … + 4 h zweiter Einsatz am selben Tag: keine Ruhezeit-Frage, 9 h gesamt
        ("2026-06-03T06:00", "2026-06-03T14:00"),   # Mi: nur 9 h Ruhe seit 21:00 → Hinweis
        ("2026-06-04T16:00", "2026-06-05T01:00"),   # Do→Fr: Dienst über Mitternacht (9 h)
        ("2026-06-05T08:00", "2026-06-05T12:00"),   # Fr: 7 h Ruhe seit 01:00 → Hinweis am Fr
    ):
        r = admin.post("/api/v1/entries", json={"user_id": uid, "start": s, "end": e})
        assert r.status_code == 201, r.text
    return uid


def days(admin, uid, f="2026-06-01", t="2026-06-05"):
    rep = admin.get("/api/v1/reports/user", params={"user_id": uid, "from": f, "to": t}).json()
    return {d["date"]: [w["code"] for w in d["warnings"]] for d in rep["days"]}, rep


def test_warnings(admin, uid):
    w, rep = days(admin, uid)
    assert w == {"2026-06-01": ["max10"], "2026-06-02": [], "2026-06-03": ["rest11"],
                 "2026-06-04": [], "2026-06-05": ["rest11"]}
    assert rep["totals"]["warnings"] == 3
    texts = {d["date"]: d["warnings"] for d in rep["days"]}
    assert texts["2026-06-03"][0]["text"] == "Ruhezeit nur 9:00 h"
    assert texts["2026-06-01"][0]["text"] == "über 10 h (11:00 h)"


def test_rest_across_month_border(admin, uid):
    # Ruhezeit-Prüfung schaut auch auf die letzte Buchung vor dem Zeitraum
    w, _ = days(admin, uid, "2026-06-03", "2026-06-03")
    assert w["2026-06-03"] == ["rest11"]


def test_in_reports(admin, uid):
    from openpyxl import load_workbook
    x = admin.get("/api/v1/reports/download", params={"user_id": uid, "period": "2026-06", "format": "xlsx"})
    notes = {r[0].day: r[12] for r in load_workbook(io.BytesIO(x.content)).active.iter_rows(values_only=True)
             if hasattr(r[0], "day")}
    assert notes[1] == "ArbZG: über 10 h (11:00 h)" and "Ruhezeit nur 9:00 h" in notes[3]
    pdf = admin.get("/api/v1/reports/download", params={"user_id": uid, "period": "2026-06", "format": "pdf"})
    assert pdf.status_code == 200


def test_switch_off(admin, uid):
    admin.patch("/api/v1/settings", json={"arbzg_warnings": False})
    try:
        w, rep = days(admin, uid)
        assert all(v == [] for v in w.values()) and rep["totals"]["warnings"] == 0
        assert admin.get("/api/v1/settings").json()["arbzg_warnings"] is False
    finally:
        admin.patch("/api/v1/settings", json={"arbzg_warnings": True})

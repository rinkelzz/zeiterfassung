"""Härtung nach dem Sicherheitsreview: CSRF, Proxy-Schlüssel, Stempel-Schlüssel, CSV-Formeln, Login-Sperre, HTTP-Hinweis."""
import io
import os

import pytest
from fastapi.testclient import TestClient

from app import csvsafe, dsm_auth
from app.main import app

LOCAL = ("127.0.0.1", 50000)


@pytest.fixture(scope="module")
def admin():
    with TestClient(app, base_url="https://zeit.example.de") as c:
        assert c.post("/api/v1/auth/login", json={"username": "admin", "password": os.environ["ZE_ADMIN_PASSWORD"]}).status_code == 200
        yield c


@pytest.fixture(scope="module")
def emp(admin):
    r = admin.post("/api/v1/users", json={"username": "harti", "display_name": "=1+2", "password": "geheim123"})
    assert r.status_code == 201, r.text
    uid = r.json()["id"]
    token = admin.post(f"/api/v1/users/{uid}/tokens", json={"label": "NFC"}).json()["token"]
    with TestClient(app, base_url="https://zeit.example.de") as c:
        assert c.post("/api/v1/auth/login", json={"username": "harti", "password": "geheim123"}).status_code == 200
        yield {"c": c, "uid": uid, "token": token}


# --- CSRF --------------------------------------------------------------------

def test_csrf_cross_site_blocked(emp):
    c = emp["c"]
    for h in ({"Sec-Fetch-Site": "cross-site"},
              {"Origin": "https://evil.example"},
              {"Origin": "null"},
              {"Sec-Fetch-Site": "same-site", "Origin": "https://blog.example.de"},
              {"Referer": "https://evil.example/seite"}):
        r = c.post("/api/v1/clock/in", content=b"", headers={"Content-Type": "text/plain", **h})
        assert r.status_code == 403, h
        assert "fremden Webseite" in r.json()["detail"]
    # Lesen bleibt erlaubt (ändert nichts)
    assert c.get("/api/v1/clock/status", headers={"Sec-Fetch-Site": "cross-site"}).status_code == 200


def test_csrf_own_page_allowed(emp):
    c = emp["c"]
    assert c.post("/api/v1/clock/in", headers={"Sec-Fetch-Site": "same-origin"}).status_code == 200
    assert c.post("/api/v1/clock/out", headers={"Origin": "https://zeit.example.de"}).status_code == 200
    assert c.post("/api/v1/clock/in").status_code == 200          # ohne Origin (Werkzeuge wie curl)
    assert c.post("/api/v1/clock/out", headers={"Origin": "https://evil.example",
                                               "X-Session-Token": "egal"}).status_code in (200, 401)  # eigener Header → kein CSRF


def test_csrf_public_url_and_forwarded_host(emp, monkeypatch):
    from app.config import settings
    c = emp["c"]
    # hinter dem Reverse Proxy kommt oft Host: localhost:8765 an – dann zählt X-Forwarded-Host bzw. die öffentliche Adresse
    h = {"Host": "localhost:8765", "Origin": "https://zeit.meine.de"}
    assert c.post("/api/v1/clock/in", headers=h).status_code == 403
    assert c.post("/api/v1/clock/in", headers={**h, "X-Forwarded-Host": "zeit.meine.de"}).status_code == 200
    monkeypatch.setattr(settings, "public_url", "https://zeit.meine.de")
    assert c.post("/api/v1/clock/out", headers=h).status_code == 200


# --- Proxy-Schlüssel nur von der NAS selbst -----------------------------------

def test_proxy_secret_only_from_loopback(admin):
    h = {"X-ZE-Proxy-Secret": dsm_auth.proxy_secret(), "X-DSM-User": "admin", "X-DSM-Admin": "1"}
    with TestClient(app, client=("8.8.8.8", 5000)) as c:
        assert c.get("/api/v1/me", headers=h).status_code == 401
    with TestClient(app, client=("192.168.1.20", 5000)) as c:
        assert c.get("/api/v1/me", headers=h).status_code == 401
    with TestClient(app, client=LOCAL) as c:
        assert c.get("/api/v1/me", headers=h).json()["via_dsm"] is True
        # über den Reverse Proxy durchgereicht (auch 127.0.0.1, aber mit X-Forwarded-For) → zählt nicht
        assert c.get("/api/v1/me", headers={**h, "X-Forwarded-For": "8.8.8.8"}).status_code == 401
        # falscher Schlüssel
        assert c.get("/api/v1/me", headers={**h, "X-ZE-Proxy-Secret": "x" * 43}).status_code == 401


# --- Stempel-Schlüssel dürfen nur stempeln ------------------------------------

def test_token_scope(admin, emp):
    with TestClient(app) as c:
        h = {"Authorization": "Bearer " + emp["token"]}
        assert c.get("/api/v1/clock/status", headers=h).status_code == 200
        assert c.post("/api/v1/clock/toggle", headers=h, json={"source": "nfc"}).status_code == 200
        assert c.post("/api/v1/clock/toggle", headers=h, json={"source": "nfc"}).status_code == 200
        assert c.get("/api/v1/me", headers=h).status_code == 200          # Einrichtungsseite
        assert c.get("/api/v1/settings", headers=h).status_code == 200    # Einrichtungsseite
        for method, path in (("GET", "/api/v1/entries"), ("GET", "/api/v1/reports/user"), ("PATCH", "/api/v1/me"),
                             ("POST", "/api/v1/me/tokens"), ("GET", "/api/v1/audit"), ("POST", "/api/v1/requests")):
            r = c.request(method, path, headers=h, json={})
            assert r.status_code == 403 and "nur zum Stempeln" in r.json()["detail"], (method, path)


def test_admin_token_cannot_admin(admin):
    me = admin.get("/api/v1/me").json()
    token = admin.post(f"/api/v1/users/{me['id']}/tokens", json={"label": "Admin-NFC"}).json()["token"]
    admin.patch("/api/v1/settings", json={"smtp_host": "mail.example.de"})
    with TestClient(app) as c:
        h = {"Authorization": "Bearer " + token}
        assert c.get("/api/v1/users", headers=h).status_code == 403
        assert c.get("/api/v1/backups", headers=h).status_code == 403
        s = c.get("/api/v1/settings", headers=h).json()
        assert not [k for k in s if k.startswith("smtp_")]
    assert admin.get("/api/v1/settings").json()["smtp_host"] == "mail.example.de"


# --- CSV/Excel: keine Formeln aus Freitext ------------------------------------

def test_csvsafe_cell():
    assert csvsafe.cell('=HYPERLINK("x")') == "'=HYPERLINK(\"x\")"
    assert csvsafe.cell("+49 30 123") == "'+49 30 123"
    assert csvsafe.cell("@SUM(A1)") == "'@SUM(A1)"
    assert csvsafe.cell("-0:14") == "-0:14" and csvsafe.cell("-1,5") == "-1,5" and csvsafe.cell("-10 min") == "-10 min"
    assert csvsafe.cell("-cmd|' /C calc'!A0") == "'-cmd|' /C calc'!A0"
    assert csvsafe.cell("Urlaub") == "Urlaub" and csvsafe.cell(5) == 5 and csvsafe.cell(None) is None


def test_csv_exports_neutralized(admin, emp):
    text = admin.get("/api/v1/reports/export.csv", params={"user_id": emp["uid"]}).text
    assert "Mitarbeiter;'=1+2" in text
    log = admin.get("/api/v1/audit/export.csv").text
    assert ";=" not in log and ";+" not in log and ";@" not in log


def test_xlsx_name_is_text(admin, emp):
    from openpyxl import load_workbook
    r = admin.get("/api/v1/reports/download", params={"user_id": emp["uid"], "period": "2026-10", "format": "xlsx"})
    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(r.content)).active
    cell = ws["A2"]
    assert cell.value.startswith("=1+2") and cell.data_type == "s"


# --- Login-Sperre: Speicher bleibt begrenzt -----------------------------------

def test_fail_table_bounded(monkeypatch):
    monkeypatch.setattr(dsm_auth, "MAX_KEYS", 50)
    dsm_auth._fails.clear()
    try:
        dsm_auth.note_failure("opfer")
        for i in range(300):
            dsm_auth.note_failure(f"name{i}")
        assert len(dsm_auth._fails) <= 51
        # abgelaufene Einträge fliegen zuerst raus
        dsm_auth._fails.clear()
        dsm_auth._fails["alt"] = [0.0]
        for i in range(60):
            dsm_auth.note_failure(f"n{i}")
        assert "alt" not in dsm_auth._fails
    finally:
        dsm_auth._fails.clear()


# --- Hinweis bei http aus dem Internet ----------------------------------------

def test_insecure_http_flag():
    with TestClient(app, client=("8.8.8.8", 5000)) as c:
        c.post("/api/v1/auth/login", json={"username": "admin", "password": os.environ["ZE_ADMIN_PASSWORD"]})
        assert c.get("/api/v1/me").json()["insecure_http"] is True
    for ip in ("192.168.178.20", "10.0.0.5", "100.101.102.103"):
        with TestClient(app, client=(ip, 5000)) as c:
            c.post("/api/v1/auth/login", json={"username": "admin", "password": os.environ["ZE_ADMIN_PASSWORD"]})
            assert c.get("/api/v1/me").json()["insecure_http"] is False, ip
    with TestClient(app, base_url="https://zeit.example.de", client=("8.8.8.8", 5000)) as c:
        c.post("/api/v1/auth/login", json={"username": "admin", "password": os.environ["ZE_ADMIN_PASSWORD"]})
        assert c.get("/api/v1/me").json()["insecure_http"] is False

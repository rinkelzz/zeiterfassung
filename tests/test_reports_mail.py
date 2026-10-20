"""Monats-/Jahresbericht als PDF/Excel, Mailversand und Zeitplan."""
import base64
import email
import email.policy
import io
import os
import socketserver
import threading
from datetime import datetime, timezone

import pytest
from fastapi.testclient import TestClient

from app import mailer
from app.db import connect
from app.main import app


# --- Mini-SMTP-Server für die Tests --------------------------------------------------

class _Handler(socketserver.StreamRequestHandler):
    def w(self, line):
        self.wfile.write((line + "\r\n").encode())

    def handle(self):
        srv = self.server
        self.w("220 test ESMTP")
        mail = {"rcpt": []}
        while True:
            line = self.rfile.readline().decode(errors="replace").rstrip("\r\n")
            if not line:
                return
            cmd = line.upper()
            if cmd.startswith("EHLO") or cmd.startswith("HELO"):
                self.w("250-test")
                self.w("250 AUTH PLAIN LOGIN")
            elif cmd.startswith("AUTH PLAIN"):
                creds = base64.b64decode(line.split()[-1]).split(b"\0")
                if creds[1:] == [b"zeit@example.de", b"richtig"]:
                    self.w("235 ok")
                else:
                    self.w("535 nope")
            elif cmd.startswith("MAIL FROM"):
                mail["from"] = line
                self.w("250 ok")
            elif cmd.startswith("RCPT TO"):
                if srv.fail_rcpt:
                    self.w("550 nicht erlaubt")
                else:
                    mail["rcpt"].append(line)
                    self.w("250 ok")
            elif cmd == "DATA":
                self.w("354 go")
                data = []
                while True:
                    ln = self.rfile.readline()
                    if ln in (b".\r\n", b".\n", b""):
                        break
                    data.append(ln[1:] if ln.startswith(b"..") else ln)
                mail["msg"] = email.message_from_bytes(b"".join(data), policy=email.policy.default)
                srv.inbox.append(mail)
                mail = {"rcpt": []}
                self.w("250 queued")
            elif cmd == "QUIT":
                self.w("221 bye")
                return
            else:
                self.w("250 ok")


class FakeSMTP(socketserver.ThreadingTCPServer):
    allow_reuse_address = True
    daemon_threads = True

    def __init__(self):
        super().__init__(("127.0.0.1", 0), _Handler)
        self.inbox = []
        self.fail_rcpt = False


@pytest.fixture(scope="module")
def smtp():
    srv = FakeSMTP()
    t = threading.Thread(target=srv.serve_forever, daemon=True)
    t.start()
    yield srv
    srv.shutdown()


@pytest.fixture(scope="module")
def admin(smtp):
    with TestClient(app) as c:
        assert c.post("/api/v1/auth/login", json={"username": "admin", "password": os.environ["ZE_ADMIN_PASSWORD"]}).status_code == 200
        r = c.patch("/api/v1/settings", json={
            "smtp_host": "127.0.0.1", "smtp_port": smtp.server_address[1], "smtp_security": "none",
            "smtp_user": "zeit@example.de", "smtp_password": "richtig", "smtp_from": "zeit@example.de",
            "smtp_from_name": "Zeiterfassung Wache", "mail_formats": ["pdf", "xlsx"], "break_mode": "none"})
        assert r.status_code == 200, r.text
        yield c


@pytest.fixture(scope="module")
def rita(admin):
    r = admin.post("/api/v1/users", json={"username": "rita", "display_name": "Rita Ränder", "password": "geheim123",
                                         "tracking_start": "2026-01-01", "weekly_hours": 35, "email": "rita@example.de"})
    assert r.status_code == 201, r.text
    uid = r.json()["id"]
    # Mo 07.09.2026: 8 h → +1 h, Ü-Tag Sa 12.09.: 4 h
    admin.post("/api/v1/entries", json={"user_id": uid, "start": "2026-09-07T07:00", "end": "2026-09-07T15:00"})
    admin.post("/api/v1/entries", json={"user_id": uid, "start": "2026-09-12T08:00", "end": "2026-09-12T12:00"})
    admin.put("/api/v1/dayflags", json={"user_id": uid, "day": "2026-09-12", "overtime": True})
    admin.post("/api/v1/absences", json={"user_id": uid, "from_day": "2026-09-14", "to_day": "2026-09-15", "kind": "urlaub"})
    return uid


def test_download_month_pdf_and_xlsx(admin, rita):
    r = admin.get("/api/v1/reports/download", params={"user_id": rita, "period": "2026-09", "format": "pdf"})
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF-1.4") and r.content.rstrip().endswith(b"%%EOF")
    assert 'filename="Arbeitszeitnachweis_rita_2026-09.pdf"' in r.headers["content-disposition"]

    from openpyxl import load_workbook
    r = admin.get("/api/v1/reports/download", params={"user_id": rita, "period": "2026-09", "format": "xlsx"})
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    ws = load_workbook(io.BytesIO(r.content)).active
    rows = {row[0].strftime("%d.%m.") if hasattr(row[0], "strftime") else row[0]: row for row in ws.iter_rows(values_only=True)}
    assert rows["07.09."][5] == 8 and rows["07.09."][7] == 1          # Ist 8 h, Diff +1 h
    assert rows["12.09."][9] == 4                                      # Ü-Tag → Überstunden
    assert rows["14.09."][11] == "Urlaub"


def test_download_year(admin, rita):
    r = admin.get("/api/v1/reports/download", params={"user_id": rita, "period": "2026", "format": "xlsx"})
    assert r.status_code == 200
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[0] == "Jahresübersicht" and "September" in wb.sheetnames and len(wb.sheetnames) == 13
    sep = [row for row in wb["Jahresübersicht"].iter_rows(values_only=True) if row[0] == "September"][0]
    assert sep[2] == 12 and sep[5] == 4                                # Ist 12 h, davon 4 h Überstunden
    r = admin.get("/api/v1/reports/download", params={"user_id": rita, "period": "2026", "format": "pdf"})
    assert r.status_code == 200 and r.content.startswith(b"%PDF")


def test_download_validation_and_rights(admin, rita):
    assert admin.get("/api/v1/reports/download", params={"period": "2026-13", "format": "pdf"}).status_code == 400
    assert admin.get("/api/v1/reports/download", params={"period": "2026-09", "format": "doc"}).status_code == 400
    with TestClient(app) as c:
        c.post("/api/v1/auth/login", json={"username": "rita", "password": "geheim123"})
        assert c.get("/api/v1/reports/download", params={"period": "2026-09"}).status_code == 200  # eigener Bericht
        assert c.get("/api/v1/reports/download", params={"period": "2026-09", "user_id": 1}).status_code == 403
        # Mailserver-Zugang sehen nur Admins
        assert "smtp_host" not in c.get("/api/v1/settings").json()


def test_settings_hide_password(admin):
    s = admin.get("/api/v1/settings").json()
    assert s["smtp_password_set"] is True and "smtp_password" not in s
    # leeres Passwort = unverändert
    admin.patch("/api/v1/settings", json={"smtp_password": ""})
    with connect() as conn:
        assert mailer.smtp_config(conn)["password"] == "richtig"
    log = admin.get("/api/v1/audit").json()
    assert all("richtig" not in (r["details"] or "") for r in log)
    assert admin.patch("/api/v1/settings", json={"smtp_from": "kaputt"}).status_code == 400
    assert admin.patch("/api/v1/settings", json={"mail_formats": []}).status_code == 400


def test_mail_test_and_report(admin, rita, smtp):
    smtp.inbox.clear()
    r = admin.post("/api/v1/settings/mail-test", json={"to": "chef@example.de"})
    assert r.status_code == 200, r.text
    assert smtp.inbox[-1]["msg"]["Subject"] == "Zeiterfassung: Testmail"

    r = admin.post("/api/v1/reports/mail", json={"user_id": rita, "period": "2026-09"})
    assert r.status_code == 200, r.text
    msg = smtp.inbox[-1]["msg"]
    assert msg["To"] == "rita@example.de" and "September 2026" in msg["Subject"]
    names = [p.get_filename() for p in msg.walk() if p.get_filename()]
    assert names == ["Arbeitszeitnachweis_rita_2026-09.pdf", "Arbeitszeitnachweis_rita_2026-09.xlsx"]
    body = [p for p in msg.walk() if p.get_content_type() == "text/plain"][0].get_content()
    assert "Ist:        12:00 h" in body and "Rita Ränder" in body

    # Admin ohne Adresse → verständlicher Fehler
    r = admin.post("/api/v1/reports/mail", json={"period": "2026-09"})
    assert r.status_code == 502 and "keine E-Mail-Adresse" in r.json()["detail"]


def test_wrong_smtp_password(admin, smtp):
    admin.patch("/api/v1/settings", json={"smtp_password": "falsch"})
    try:
        r = admin.post("/api/v1/settings/mail-test", json={"to": "chef@example.de"})
        assert r.status_code == 502 and "Anmeldung am Mailserver" in r.json()["detail"]
    finally:
        admin.patch("/api/v1/settings", json={"smtp_password": "richtig"})


def test_profile_email(admin):
    with TestClient(app) as c:
        admin.post("/api/v1/users", json={"username": "paul", "display_name": "Paul", "password": "geheim123"})
        c.post("/api/v1/auth/login", json={"username": "paul", "password": "geheim123"})
        assert c.patch("/api/v1/me", json={"email": "kein-mail"}).status_code == 400
        me = c.patch("/api/v1/me", json={"email": " paul@example.de ", "report_mail": False}).json()
        assert me["email"] == "paul@example.de" and me["report_mail"] is False
    assert admin.post("/api/v1/users", json={"username": "x1", "display_name": "X", "password": "geheim123",
                                            "email": "falsch"}).status_code == 422


def _at(y, m, d, h, mi=0):
    """Ortszeit Berlin → UTC (Oktober/Januar: +2/+1 Stunden)."""
    from app.timeutil import tz
    return datetime(y, m, d, h, mi, tzinfo=tz()).astimezone(timezone.utc)


def test_scheduler_monthly(admin, rita, smtp):
    admin.patch("/api/v1/settings", json={"mail_monthly": True, "mail_yearly": False})
    smtp.inbox.clear()
    with connect() as conn:
        conn.execute("DELETE FROM report_mails")
        assert mailer.run_due(conn, _at(2026, 10, 1, 5, 30)) == []         # vor 6 Uhr nichts
        res = mailer.run_due(conn, _at(2026, 10, 1, 6, 5))
        sent_to = [m["msg"]["To"] for m in smtp.inbox]
        assert "rita@example.de" in sent_to and "paul@example.de" not in sent_to  # Paul will keine Mails
        assert all(r["period"] == "2026-09" for r in res)
        assert mailer.run_due(conn, _at(2026, 10, 1, 7)) == []             # nur einmal
        conn.execute("DELETE FROM report_mails")
        assert mailer.run_due(conn, _at(2026, 10, 8, 7)) == []             # nach dem 7. nicht nachholen
        assert len(mailer.run_due(conn, _at(2026, 10, 5, 12))) >= 1        # bis zum 7. nachholen


def test_scheduler_retry_and_yearly(admin, rita, smtp):
    admin.patch("/api/v1/settings", json={"mail_monthly": False, "mail_yearly": True})
    with connect() as conn:
        conn.execute("DELETE FROM report_mails")
        smtp.fail_rcpt = True
        try:
            res = mailer.run_due(conn, _at(2027, 1, 1, 6, 10))
            assert res and not res[0]["ok"] and res[0]["period"] == "2026"
            # sofort danach kein neuer Versuch, nach einer Stunde schon
            assert mailer.run_due(conn, _at(2027, 1, 1, 6, 40)) == []
        finally:
            smtp.fail_rcpt = False
        res = mailer.run_due(conn, _at(2027, 1, 1, 7, 20))
        assert res and res[0]["ok"]
        log = conn.execute("SELECT * FROM report_mails WHERE user_id = ? AND period = '2026'", (rita,)).fetchone()
        assert log["status"] == "sent" and log["attempts"] == 2
    assert any(r["period"] == "2026" for r in admin.get("/api/v1/reports/mail-log").json())
    admin.patch("/api/v1/settings", json={"mail_yearly": False})

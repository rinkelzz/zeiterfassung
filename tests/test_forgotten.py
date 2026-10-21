"""Vergessenes Ausstempeln: Erinnerung nach 12 h, automatisch beenden um Mitternacht, „bitte prüfen“."""
import os
import threading
from datetime import datetime

import pytest
from fastapi.testclient import TestClient

from app import forgotten
from app.db import connect
from app.main import app
from app.timeutil import tz
from test_reports_mail import FakeSMTP


def at(y, m, d, h, mi=0):
    return datetime(y, m, d, h, mi, tzinfo=tz())


@pytest.fixture(scope="module")
def smtp():
    srv = FakeSMTP()
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield srv
    srv.shutdown()


@pytest.fixture(scope="module")
def admin(smtp):
    with TestClient(app) as c:
        assert c.post("/api/v1/auth/login", json={"username": "admin", "password": os.environ["ZE_ADMIN_PASSWORD"]}).status_code == 200
        c.patch("/api/v1/settings", json={
            "smtp_host": "127.0.0.1", "smtp_port": smtp.server_address[1], "smtp_security": "none",
            "smtp_user": "zeit@example.de", "smtp_password": "richtig", "smtp_from": "zeit@example.de",
            "open_limit_hours": 12, "open_auto_close": True, "break_mode": "none"})
        c.patch("/api/v1/users/1", json={"email": "chef@example.de"})
        yield c
        c.patch("/api/v1/users/1", json={"email": None})  # andere Tests erwarten den Admin ohne Adresse


def test_settings_roundtrip(admin):
    s = admin.patch("/api/v1/settings", json={"open_limit_hours": 10.5}).json()
    assert s["open_limit_hours"] == 10.5 and s["open_auto_close"] is True
    assert admin.patch("/api/v1/settings", json={"open_limit_hours": 99}).status_code == 422
    admin.patch("/api/v1/settings", json={"open_limit_hours": 12})


def test_close_time():
    from app.timeutil import to_local
    start = at(2026, 9, 7, 6, 50).astimezone(tz())
    assert to_local(forgotten.close_time(start, 12)) == at(2026, 9, 8, 0, 0)    # Grenze 18:50 → Mitternacht
    start = at(2026, 9, 7, 20, 0)
    assert to_local(forgotten.close_time(start, 12)) == at(2026, 9, 9, 0, 0)    # Grenze 08:00 am 8. → Mitternacht danach


def test_remind_then_close(admin, smtp):
    r = admin.post("/api/v1/users", json={"username": "vergess", "display_name": "Vera Vergess", "password": "geheim123",
                                         "tracking_start": "2026-01-01", "email": "vera@example.de"})
    uid = r.json()["id"]
    r = admin.post("/api/v1/entries", json={"user_id": uid, "start": "2026-09-07T06:50"})
    assert r.status_code == 201, r.text
    eid = r.json()["id"]
    smtp.inbox.clear()
    with connect() as conn:
        assert forgotten.run(conn, at(2026, 9, 7, 18, 0)) == {"reminded": [], "closed": []}   # erst 11 h
        res = forgotten.run(conn, at(2026, 9, 7, 19, 0))
        assert eid in res["reminded"] and not res["closed"]
        to = sorted(m["msg"]["To"] for m in smtp.inbox)
        assert to == ["chef@example.de", "vera@example.de"]
        body = smtp.inbox[0]["msg"].get_body(("plain",)).get_content()
        assert "seit 07.09.2026 um 06:50 Uhr" in body and "07.09. 18:50 Uhr" in body
        assert forgotten.run(conn, at(2026, 9, 7, 22, 0))["reminded"] == []                  # nur einmal
        assert len(smtp.inbox) == 2
        res = forgotten.run(conn, at(2026, 9, 8, 0, 5))
        assert res["closed"] == [eid]
        e = conn.execute("SELECT * FROM entries WHERE id = ?", (eid,)).fetchone()
        assert e["auto_closed"] == 1 and "bitte prüfen" in e["note"]
    day = admin.get("/api/v1/reports/user", params={"user_id": uid, "from": "2026-09-07", "to": "2026-09-07"}).json()["days"][0]
    assert day["needs_check"] and day["last_out"] == "18:50" and day["worked"] == "12:00"
    # PDF/Excel enthalten den Hinweis
    from openpyxl import load_workbook
    import io
    x = admin.get("/api/v1/reports/download", params={"user_id": uid, "period": "2026-09", "format": "xlsx"})
    rows = [r for r in load_workbook(io.BytesIO(x.content)).active.iter_rows(values_only=True) if r[1] == "Mo" and r[2] == "06:50"]
    assert "bitte prüfen" in rows[0][12]
    # Admin korrigiert → Markierung weg
    r = admin.patch(f"/api/v1/entries/{eid}", json={"end": "2026-09-07T15:30"})
    assert r.status_code == 200 and r.json()["auto_closed"] is False
    day = admin.get("/api/v1/reports/user", params={"user_id": uid, "from": "2026-09-07", "to": "2026-09-07"}).json()["days"][0]
    assert not day["needs_check"] and day["last_out"] == "15:30"


def test_no_auto_close_and_off(admin):
    r = admin.post("/api/v1/users", json={"username": "lang", "display_name": "Lang", "password": "geheim123"})
    uid = r.json()["id"]
    eid = admin.post("/api/v1/entries", json={"user_id": uid, "start": "2026-09-10T07:00"}).json()["id"]
    admin.patch("/api/v1/settings", json={"open_auto_close": False})
    try:
        with connect() as conn:
            res = forgotten.run(conn, at(2026, 9, 12, 9, 0))
            assert eid in res["reminded"] and not res["closed"]   # nur erinnert (keine Adresse → keine Mail)
            admin.patch("/api/v1/settings", json={"open_limit_hours": 0})
            assert forgotten.run(conn, at(2026, 9, 20, 9, 0)) == {"reminded": [], "closed": []}
            assert conn.execute("SELECT end_utc FROM entries WHERE id = ?", (eid,)).fetchone()[0] is None
    finally:
        admin.patch("/api/v1/settings", json={"open_limit_hours": 12, "open_auto_close": True})
        admin.delete(f"/api/v1/entries/{eid}")
